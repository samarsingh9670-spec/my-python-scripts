import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import tkinter as tk
from tkinter import ttk
from tkinter import messagebox
from tkcalendar import DateEntry

class ImplementationPlanGUI:

    def __init__(self, root):

        self.root = root
        self.root.title("Implementation Plan Generator")
        self.root.geometry("900x600")

        self.batch_entries = []

        self.build_gui()

    # ==================================================
    # BUILD GUI
    # ==================================================

    def build_gui(self):

        main = ttk.Frame(self.root, padding=15)
        main.pack(fill="both", expand=True)

        ttk.Label(
            main,
            text="Implementation Plan Generator",
            font=("Segoe UI", 16, "bold")
        ).grid(
            row=0,
            column=0,
            columnspan=4,
            pady=15
        )

        # ==========================================
        # Platform
        # ==========================================

        ttk.Label(
            main,
            text="Platform"
        ).grid(
            row=1,
            column=0,
            sticky="w",
            pady=5
        )

        self.platform = tk.StringVar()

        platform_combo = ttk.Combobox(
            main,
            textvariable=self.platform,
            values=["AMC", "AVCC"],
            state="readonly",
            width=30
        )

        platform_combo.grid(
            row=1,
            column=1,
            sticky="w"
        )

        platform_combo.current(0)

        # ==========================================
        # New Request
        # ==========================================

        ttk.Label(
            main,
            text="New Request"
        ).grid(
            row=2,
            column=0,
            sticky="w",
            pady=5
        )

        self.new_request = tk.StringVar(value="YES")

        ttk.Radiobutton(
            main,
            text="Yes",
            variable=self.new_request,
            value="YES",
            command=self.toggle_previous_section
        ).grid(
            row=2,
            column=1,
            sticky="w"
        )

        ttk.Radiobutton(
            main,
            text="No",
            variable=self.new_request,
            value="NO",
            command=self.toggle_previous_section
        ).grid(
            row=2,
            column=2,
            sticky="w"
        )

        # ==========================================
        # Application Details
        # ==========================================

        self.app_id = tk.StringVar()
        self.app_name = tk.StringVar()
        self.app_link = tk.StringVar()
        self.prod_group = tk.StringVar()

        self.create_entry(
            main,
            "Application ID",
            self.app_id,
            3
        )

        self.create_entry(
            main,
            "Application Name",
            self.app_name,
            4
        )

        self.create_entry(
            main,
            "Application Intune Link",
            self.app_link,
            5
        )

        self.create_entry(
            main,
            "Production Group",
            self.prod_group,
            6
        )

        # ==========================================
        # Previous Version Section
        # ==========================================

        self.previous_frame = ttk.LabelFrame(
            main,
            text="Previous Version Details"
        )

        self.previous_frame.grid(
            row=7,
            column=0,
            columnspan=4,
            sticky="ew",
            pady=10
        )

        self.previous_app_name = tk.StringVar()
        self.previous_app_link = tk.StringVar()

        ttk.Label(
            self.previous_frame,
            text="Previous Application Name"
        ).grid(
            row=0,
            column=0,
            padx=5,
            pady=5
        )

        ttk.Entry(
            self.previous_frame,
            textvariable=self.previous_app_name,
            width=55
        ).grid(
            row=0,
            column=1,
            padx=5
        )

        ttk.Label(
            self.previous_frame,
            text="Previous Application Link"
        ).grid(
            row=1,
            column=0,
            padx=5,
            pady=5
        )

        ttk.Entry(
            self.previous_frame,
            textvariable=self.previous_app_link,
            width=55
        ).grid(
            row=1,
            column=1,
            padx=5
        )

        self.previous_frame.grid_remove()

        # ==========================================
        # Batch Count
        # ==========================================

        ttk.Label(
            main,
            text="Number of Deployment Batches"
        ).grid(
            row=8,
            column=0,
            sticky="w",
            pady=10
        )

        self.batch_count = tk.StringVar()

        ttk.Entry(
            main,
            textvariable=self.batch_count,
            width=10
        ).grid(
            row=8,
            column=1,
            sticky="w"
        )

        ttk.Button(
            main,
            text="Create Batch Fields",
            command=self.create_batch_fields
        ).grid(
            row=8,
            column=2,
            padx=10
        )

        # ==========================================
        # Batch Frame
        # ==========================================

        self.batch_frame = ttk.LabelFrame(
            main,
            text="Deployment Batches"
        )

        self.batch_frame.grid(
            row=9,
            column=0,
            columnspan=4,
            sticky="ew",
            pady=15
        )

        # ==========================================
        # Generate Button
        # ==========================================

        ttk.Button(
            main,
            text="Generate Document",
            command=self.generate_data
        ).grid(
            row=10,
            column=0,
            columnspan=4,
            pady=20
        )

    # ==================================================
    # Create Entry
    # ==================================================

    def create_entry(
        self,
        parent,
        label,
        variable,
        row
    ):

        ttk.Label(
            parent,
            text=label
        ).grid(
            row=row,
            column=0,
            sticky="w",
            pady=5
        )

        ttk.Entry(
            parent,
            textvariable=variable,
            width=70
        ).grid(
            row=row,
            column=1,
            columnspan=3,
            sticky="w"
        )

    # ==================================================
    # Show / Hide Previous Version
    # ==================================================

    def toggle_previous_section(self):

        if self.new_request.get() == "NO":
            self.previous_frame.grid()
        else:
            self.previous_frame.grid_remove()

    # ==================================================
    # Dynamic Batch Creation
    # ==================================================

    def create_batch_fields(self):

        for widget in self.batch_frame.winfo_children():
            widget.destroy()

        self.batch_entries.clear()

        try:
            count = int(self.batch_count.get())

            if count <= 0:
                raise ValueError

        except ValueError:

            messagebox.showerror(
                "Error",
                "Enter a valid batch number."
            )
            return

        for i in range(count):

            ttk.Label(
                self.batch_frame,
                text=f"Batch {i+1} Date"
            ).grid(
                row=i,
                column=0,
                padx=5,
                pady=5
            )

            date_picker = DateEntry(
                self.batch_frame,
                width=18,
                date_pattern="dd-mm-yyyy"
            )

            date_picker.grid(
                row=i,
                column=1,
                padx=5
            )

            ttk.Label(
                self.batch_frame,
                text="User Count"
            ).grid(
                row=i,
                column=2,
                padx=5
            )

            user_var = tk.StringVar()

            ttk.Entry(
                self.batch_frame,
                textvariable=user_var,
                width=15
            ).grid(
                row=i,
                column=3,
                padx=5
            )

            self.batch_entries.append(
                (
                    date_picker,
                    user_var
                )
            )

    # ==================================================
    # Generate Data
    # ==================================================

    def generate_data(self):

        batches = []

        for i, (date_picker, user_var) in enumerate(
            self.batch_entries,
            start=1
        ):

            batches.append({
                "batch": f"Batch {i}",
                "planned_date": date_picker.get(),
                "user_count": user_var.get()
            })
        data = {

            "platform":
                self.platform.get(),

            "new_request":
                self.new_request.get(),

            "application_id":
                self.app_id.get(),

            "application_name":
                self.app_name.get(),

            "application_link":
                self.app_link.get(),

            "production_group":
                self.prod_group.get(),

            "previous_application_name":
                self.previous_app_name.get(),

            "previous_application_link":
                self.previous_app_link.get(),

            "batches":
                batches
        }

        print("\n===== COLLECTED DATA =====")
        print(data)

        create_plan(data)

#===================================
#DOCX Function
#===================================
def add_bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")

from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_hyperlink(paragraph, text, url):

    part = paragraph.part

    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")

    rPr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")

    rPr.append(color)
    rPr.append(underline)

    new_run.append(rPr)

    text_element = OxmlElement("w:t")
    text_element.text = text

    new_run.append(text_element)
    hyperlink.append(new_run)

    paragraph._p.append(hyperlink)

    return hyperlink

def create_plan(data):
    print("create_plan started")
    platform = data["platform"]

    if platform == "AMC":
        filter_name = "ALL AMC Devices"
        output_suffix = "AMC"
    else:
        filter_name = "AVC - All PROD CloudPCs"
        output_suffix = "AVCC"

    new_request = data["new_request"]

    app_id = data["application_id"]
    app_name = data["application_name"]
    app_link = data["application_link"]
    prod_group = data["production_group"]

    previous_app_name = data["previous_application_name"]
    previous_app_link = data["previous_application_link"]

    batches = data["batches"]

    document = Document()

    document.add_heading(
        "Standard Implementation and Backout Plan for Intune Application Deployments",
        level=0
    )

    document.add_heading(
        "Application Details",
        level=1
    )

    add_bullet(
        document,
        f"Platform: {platform}"
    )

    add_bullet(
        document,
        f"Application ID: {app_id}"
    )

    add_bullet(
        document,
        f"Application Name: {app_name}"
    )

    p = document.add_paragraph(
        style="List Bullet"
    )

    p.add_run(
        "Application Intune Link: "
    )

    add_hyperlink(
        p,
        app_name,
        app_link
    )

    # Deployment Schedule

    document.add_heading(
        "Deployment Schedule",
        level=2
    )

    table = document.add_table(
        rows=1,
        cols=3
    )

    table.style = "Table Grid"

    hdr = table.rows[0].cells

    hdr[0].text = "Batch"
    hdr[1].text = "Planned Date"
    hdr[2].text = "User Count"

    for batch in batches:

        row = table.add_row().cells

        row[0].text = batch["batch"]
        row[1].text = batch["planned_date"]
        row[2].text = batch["user_count"]

    # Implementation Steps

    document.add_heading(
        "Implementation Steps",
        level=2
    )

    implementation_steps = [

        "Log in to the Microsoft Intune Admin Center.",

        f"Navigate to Apps → Windows (By Platform) → {app_name} → Properties.",

        "Edit the Assignments section.",

        f"Add the required Batchwise Group or {prod_group} with assignment filter {filter_name}.",

        "Click Review + Save and verify the configuration.",

        "Validate successful deployment for each batch before proceeding to the next batch."

    ]

    if new_request == "NO":

        implementation_steps.append(
            f"Upon successful deployment to Production Group, remove the assignments from the previous application version."
        )

    for step in implementation_steps:
        add_bullet(document, step)

    # Previous Version

    if new_request == "NO":

        document.add_heading(
            "Previous Version Details",
            level=2
        )

        add_bullet(
            document,
            f"Application Name: {previous_app_name}"
        )

        add_bullet(
            document,
            f"Assignments to be Removed: {prod_group}"
        )

        p2 = document.add_paragraph(
            style="List Bullet"
        )

        p2.add_run(
            "Application Intune Link: "
        )

    # Previous Version Assignment Removal Steps

    if new_request == "NO":

        document.add_heading(
        "Previous Version Assignment Removal Steps",
        level=2
        )

        add_bullet(
            document,
            f"Log in to the Microsoft Intune Admin Center."
        )

        add_bullet(
            document,
            f"Navigate to Apps → Windows (By Platform) → {previous_app_name} → Properties."
        )

        add_bullet(
            document,
            f"Edit the Assignments section."
        )

        add_bullet(
            document,
            f"Remove application Production Group {prod_group}."
        )

        add_bullet(
            document,
            f"Click Review + Save and verify the configuration."
        )

        add_hyperlink(
            p2,
            previous_app_name,
            previous_app_link
        )
    
    # Backout Plan

    document.add_heading(
        "Backout Plan",
        level=2
    )

    add_bullet(
        document,
        "Log in to the Microsoft Intune Admin Center."
    )

    add_bullet(
        document,
        f"Navigate to Apps → Windows (By Platform) → {app_name} → Properties."
    )

    add_bullet(
        document,
        f"Remove Production Group {prod_group} or Batchwise Groups."
    )

    add_bullet(
        document,
        f"Click Review + Save."
    )

    add_bullet(
        document,
        f"If assignments were removed from the previous version, reassign the original groups to restore the previous deployment state."
    )

    save_folder = rf"C:\Users\HCL0190\OneDrive - Allianz\Desktop\App Artifacts\{app_id}"

    os.makedirs(
        save_folder,
        exist_ok=True
    )

    file_path = os.path.join(
        save_folder,
        f"{app_id} {output_suffix} Implementation Plan.docx"
    )

    document.save(
        file_path
    )

    messagebox.showinfo(
        "Success",
        f"Document created successfully\n\n{file_path}"
    )

# ==================================================
# START APPLICATION
# ==================================================

if __name__ == "__main__":

    root = tk.Tk()

    app = ImplementationPlanGUI(root)

    root.mainloop()